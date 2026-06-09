"""Build stock.xlsx end-to-end from scratch — items, suppliers, transactions, styling.

Writes the seed workbook students start from to references/stock.xlsx (next to
this script). The completed model solution lives alongside it as
references/jugaad_inventory.ods.
"""
import random
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT_PATH = str(Path(__file__).resolve().parent / "references" / "stock.xlsx")

# ─── Workbook scaffold ──────────────────────────────────────────────────────
wb = Workbook()
wb.remove(wb.active)
for name in ['Dashboard', 'Item_Master', 'Suppliers', 'Stock_Transactions', 'Current_Stock']:
    wb.create_sheet(name)

# ─── Suppliers (seed 42) ────────────────────────────────────────────────────
random.seed(42)

SUPPLIER_NAMES = [
    "Modern Enterprises", "Diamond Sales Corp.", "Mahalaxmi Trading House",
    "Golden Enterprises", "Imperial Enterprises", "Premier Sales Corp.",
    "Diamond Trading Co.", "Galaxy Trading House", "Northern Distributors",
    "Eastern Distributors", "New Enterprises", "Galaxy Agencies",
    "Mahalaxmi Tools & Hardware", "Sunrise Enterprises", "Bharat Trading House",
    "Capital Agencies", "Western Tools & Hardware", "Southern Enterprises",
    "Southern Sales Corp.", "Southern Distributors", "National Trading Co.",
    "New Trading Co.", "Universal Marketing", "Apex Hardware Mart",
    "Reliance Tools Hub", "Metro Industrial Supply", "Skyline Traders",
    "Heritage Hardware Co.",
]
CITIES = [
    ("Indore", "Madhya Pradesh"), ("Jaipur", "Rajasthan"), ("Kolkata", "West Bengal"),
    ("Nagpur", "Maharashtra"), ("Hyderabad", "Telangana"), ("Lucknow", "Uttar Pradesh"),
    ("Pune", "Maharashtra"), ("Coimbatore", "Tamil Nadu"), ("Chennai", "Tamil Nadu"),
    ("Mumbai", "Maharashtra"), ("Gurgaon", "Haryana"), ("Kochi", "Kerala"),
    ("Mysore", "Karnataka"), ("Faridabad", "Haryana"), ("Delhi", "Delhi"),
    ("Bengaluru", "Karnataka"), ("Ahmedabad", "Gujarat"), ("Surat", "Gujarat"),
]
FIRST = ["Mohan", "Sunil", "Anil", "Rakesh", "Ashish", "Sandeep", "Meera",
         "Pradeep", "Ramesh", "Manoj", "Deepak", "Kiran", "Anjali", "Vikram",
         "Rajesh", "Amit", "Kavita", "Neha", "Pooja", "Suresh"]
LAST = ["Patel", "Sharma", "Khan", "Reddy", "Gupta", "Iyer", "Singh",
        "Verma", "Joshi", "Nair", "Mehta", "Rao"]
TERMS = ["COD", "Net 15", "Net 30", "Net 45", "Net 60", "50% Advance"]

suppliers = []
for i, name in enumerate(SUPPLIER_NAMES, start=1):
    city, state = random.choice(CITIES)
    suppliers.append({
        "Supplier ID": f"SUP-{i:02d}",
        "Supplier Name": name,
        "City": city,
        "State": state,
        "Contact Person": f"{random.choice(FIRST)} {random.choice(LAST)}",
        "Phone": f"{random.randint(6, 9)}{random.randint(100000000, 999999999)}",
        "Payment Terms": random.choice(TERMS),
    })

# ─── Items (continues seed 42) ──────────────────────────────────────────────
CATALOG = {
    "Power Tools": {
        "brands": ["Bosch", "DeWalt", "Makita", "Black & Decker", "Stanley", "Crompton"],
        "types": [
            ("Cordless Drill", ["12V", "18V", "20V Pro"]),
            ("Hammer Drill", ["Standard", "Heavy Duty"]),
            ("Angle Grinder", ["4 inch", "5 inch", "9 inch"]),
            ("Circular Saw", ["7 inch", "8 inch", "10 inch"]),
            ("Jig Saw", ["Standard", "Variable Speed", "Pro"]),
            ("Reciprocating Saw", ["Standard", "Pro"]),
            ("Heat Gun", ["1500W", "2000W"]),
            ("Bench Grinder", ["6 inch", "8 inch"]),
            ("Rotary Hammer", ["SDS Plus", "SDS Max"]),
            ("Random Orbit Sander", ["125mm", "150mm"]),
        ],
        "cost_range": (2000, 22000),
    },
    "Hand Tools": {
        "brands": ["Taparia", "Stanley", "Visko", "Bosi", "TVS", "Pye"],
        "types": [
            ("Claw Hammer", ["8oz", "16oz", "20oz"]),
            ("Ball Peen Hammer", ["8oz", "16oz"]),
            ("Adjustable Wrench", ["6 inch", "10 inch", "12 inch"]),
            ("Pipe Wrench", ["10 inch", "14 inch"]),
            ("Pliers", ["Combination", "Long Nose", "Cutting"]),
            ("Screwdriver Set", ["6 pc", "12 pc"]),
            ("Allen Key Set", ["Metric", "Imperial"]),
            ("Spanner Set", ["8 pc", "12 pc"]),
            ("Tape Measure", ["3m", "5m", "8m"]),
            ("Spirit Level", ["12 inch", "24 inch"]),
            ("Hacksaw Frame", ["Standard"]),
            ("Socket", ["1/2 inch Drive Set"]),
        ],
        "cost_range": (100, 3500),
    },
    "Hardware": {
        "brands": ["Godrej", "Yale", "Hettich", "Hafele", "Europa"],
        "types": [
            ("Door Hinge", ["3 inch", "4 inch", "6 inch"]),
            ("Door Handle", ["Stainless", "Brass"]),
            ("Door Closer", ["Standard", "Heavy Duty"]),
            ("Door Stopper", ["Magnetic", "Floor"]),
            ("Lock", ["Mortise", "Cylinder", "Padlock"]),
            ("Cabinet Knob", ["Steel", "Chrome"]),
            ("Drawer Slide", ["12 inch", "18 inch"]),
            ("Bracket", ["L-Shape", "Heavy"]),
            ("Door Eye", ["Standard"]),
        ],
        "cost_range": (80, 2500),
    },
    "Fasteners": {
        "brands": ["GKW", "Apex", "Anchor"],
        "types": [
            ("Hex Bolts", ["M6", "M8", "M10", "M12"]),
            ("Machine Screws", ["M3", "M4", "M5"]),
            ("Wood Screws", ["1 inch", "1.5 inch", "2 inch"]),
            ("Self-tapping Screws", ["No.8", "No.10"]),
            ("Nuts", ["M6", "M8", "M10"]),
            ("Washers", ["M6", "M8", "M10"]),
            ("Anchor Bolts", ["M8", "M10"]),
            ("Wall Plugs", ["6mm", "8mm", "10mm"]),
            ("Iron Nails", ["1 inch", "2 inch", "3 inch"]),
        ],
        "cost_range": (40, 600),
        "unit_bias": "pack",
    },
    "Electrical": {
        "brands": ["Havells", "Anchor", "Plaza", "Polycab", "Legrand", "Philips", "Finolex", "Crompton"],
        "types": [
            ("LED Bulb", ["9W", "12W", "15W"]),
            ("LED Tube Light", ["18W", "22W"]),
            ("Ceiling Fan", ["48 inch", "52 inch"]),
            ("Exhaust Fan", ["6 inch", "8 inch"]),
            ("Switch", ["1-way", "2-way"]),
            ("MCB", ["6A", "16A", "32A"]),
            ("PVC Conduit", ["20mm", "25mm"]),
            ("PVC Wire", ["1.5 sqmm", "2.5 sqmm", "4 sqmm"]),
        ],
        "cost_range": (60, 4500),
    },
    "Plumbing": {
        "brands": ["Astral", "Supreme", "Prince", "Ashirvad", "Finolex"],
        "types": [
            ("PVC Pipe", ["1/2 inch", "3/4 inch", "1 inch"]),
            ("CPVC Pipe", ["1/2 inch", "3/4 inch"]),
            ("PVC Elbow", ["1/2 inch", "1 inch"]),
            ("PVC Tee", ["1/2 inch", "1 inch"]),
            ("Ball Valve", ["1/2 inch", "1 inch"]),
            ("Bib Tap", ["Standard"]),
            ("Pillar Cock", ["Standard"]),
            ("Shower Head", ["Standard", "Rainfall"]),
            ("Towel Rod", ["18 inch", "24 inch"]),
            ("Teflon Tape", ["12mm", "19mm"]),
        ],
        "cost_range": (40, 2200),
    },
    "Paints & Coatings": {
        "brands": ["Asian Paints", "Berger", "Nerolac", "Dulux", "Birla"],
        "types": [
            ("Interior Emulsion", ["1L", "4L", "10L"]),
            ("Exterior Emulsion", ["1L", "4L", "10L"]),
            ("Enamel Paint", ["1L", "4L"]),
            ("Distemper", ["5kg", "10kg"]),
            ("Wall Putty", ["5kg", "20kg", "40kg"]),
            ("Wall Primer", ["1L", "4L"]),
            ("Metal Primer", ["1L", "4L"]),
            ("Wood Stain", ["1L"]),
            ("Wood Varnish", ["1L"]),
        ],
        "cost_range": (200, 4000),
    },
    "Adhesives & Sealants": {
        "brands": ["Pidilite", "Fevicol", "Araldite", "M-Seal", "Loctite", "3M"],
        "types": [
            ("Wood Glue", ["100g", "500g", "1kg"]),
            ("Super Glue", ["3g", "20g"]),
            ("Epoxy Resin", ["100g", "500g"]),
            ("Silicone Sealant", ["280ml"]),
            ("M-Seal", ["50g", "100g"]),
            ("Insulation Foam", ["750ml"]),
            ("Duct Tape", ["48mm x 25m"]),
            ("Masking Tape", ["24mm", "48mm"]),
            ("Insulation Tape", ["18mm"]),
            ("Double-sided Tape", ["24mm"]),
        ],
        "cost_range": (40, 1500),
    },
    "Safety Gear": {
        "brands": ["Karam", "Mallcom", "Udyogi", "Honeywell", "3M"],
        "types": [
            ("Safety Helmet", ["White", "Yellow", "Blue"]),
            ("Safety Goggles", ["Clear", "Tinted"]),
            ("Work Gloves", ["Cotton", "Leather", "Nitrile"]),
            ("Safety Boots", ["Steel Toe"]),
            ("Ear Plugs", ["Foam", "Silicone"]),
            ("Dust Mask", ["N95", "Reusable"]),
            ("High-vis Vest", ["L", "XL"]),
            ("Safety Harness", ["Full Body"]),
            ("First Aid Kit", ["Small", "Industrial"]),
        ],
        "cost_range": (80, 3500),
    },
    "Garden Tools": {
        "brands": ["Wolf-Garten", "Falcon", "Trustbasket"],
        "types": [
            ("Pruning Shears", ["Standard", "Heavy Duty"]),
            ("Spade", ["Standard"]),
            ("Shovel", ["Standard"]),
            ("Rake", ["Steel", "Plastic"]),
            ("Hoe", ["Standard"]),
            ("Watering Can", ["5L", "10L"]),
            ("Garden Hose", ["10m", "20m"]),
            ("Sprayer", ["1L", "5L"]),
            ("Lawn Mower", ["Manual", "Electric"]),
        ],
        "cost_range": (150, 8000),
    },
}

NUM_ITEMS = 350
weights = [60, 50, 35, 35, 45, 35, 30, 30, 20, 20]
plan = []
for cat, w in zip(CATALOG.keys(), weights):
    plan.extend([cat] * w)
plan = plan[:NUM_ITEMS]
random.shuffle(plan)

items = []
for idx, cat in enumerate(plan, start=1):
    info = CATALOG[cat]
    brand = random.choice(info["brands"])
    ptype, specs = random.choice(info["types"])
    spec = random.choice(specs)
    name = f"{brand} {ptype} {spec}".strip()
    lo, hi = info["cost_range"]
    cost = int(random.triangular(lo, hi, lo + (hi - lo) * 0.3) / 10) * 10
    sell = int(cost * random.uniform(1.18, 1.45) / 10) * 10
    if info.get("unit_bias") == "pack":
        unit = random.choices(["pack", "each", "unit"], weights=[6, 2, 2])[0]
    else:
        unit = random.choices(["each", "pack", "unit"], weights=[7, 1, 2])[0]
    items.append({
        "Item ID": f"HW-{idx:04d}",
        "Item Name": name,
        "Category": cat,
        "Brand": brand,
        "Unit": unit,
        "Supplier ID": random.choice(suppliers)["Supplier ID"],
        "Cost Price": cost,
        "Selling Price": sell,
        "Reorder Level": random.choice([5, 10, 15, 20, 25, 30]),
        "Opening Stock": random.randint(0, 80),
    })

# ─── Transactions (seed 7, realistic constraints) ───────────────────────────
random.seed(7)
START = date(2025, 10, 1)
END = date(2026, 3, 31)
DAYS = (END - START).days

N = len(items)
n_oos = int(N * 0.07)
n_reorder = int(N * 0.13)
n_ok = N - n_oos - n_reorder
targets = ['OOS'] * n_oos + ['REORDER'] * n_reorder + ['OK'] * n_ok
random.shuffle(targets)

NOTES = {
    "OUT":    ["Contractor account", "Walk-in customer", "Bulk order", "Site delivery", None, None, None, None],
    "IN":     ["Replenishment", "Bulk purchase", None, None, None, None],
    "DAMAGE": ["Damaged in transit", "Storage damage", "Mishandled", None, None],
    "RETURN": ["Customer return", "Wrong item delivered", "Defective", None, None],
}
def ref(t):
    return {
        "IN":     f"PO-2026-{random.randint(1000, 9999)}",
        "OUT":    f"INV-{random.randint(10000, 99999)}",
        "DAMAGE": f"DMG-{random.randint(100, 999)}",
        "RETURN": f"RET-{random.randint(1000, 9999)}",
    }[t]

all_txns = []
for it, target in zip(items, targets):
    iid, opening, reorder = it["Item ID"], it["Opening Stock"], it["Reorder Level"]
    stock = opening
    n = random.randint(6, 14)
    dates = sorted(START + timedelta(days=random.randint(0, DAYS - 14)) for _ in range(n))
    item_txns = []
    for d in dates:
        ttype = random.choices(["OUT", "IN", "DAMAGE", "RETURN"], weights=[55, 30, 8, 7])[0]
        if ttype == "IN":
            qty = random.choice([20, 30, 40, 50, 60, 80, 100])
            stock += qty
        elif ttype == "OUT":
            if stock <= 0: continue
            qty = random.randint(1, min(stock, 40))
            stock -= qty
        elif ttype == "DAMAGE":
            if stock <= 0: continue
            qty = random.randint(1, min(stock, 5))
            stock -= qty
        else:
            qty = random.randint(1, 5)
            stock += qty
        item_txns.append({"Date": d, "Item ID": iid, "Type": ttype, "Quantity": qty,
                          "Reference": ref(ttype), "Notes": random.choice(NOTES[ttype])})

    fd = END - timedelta(days=random.randint(0, 14))
    if target == 'OOS' and stock > 0:
        item_txns.append({"Date": fd, "Item ID": iid, "Type": "OUT", "Quantity": stock,
                          "Reference": ref("OUT"), "Notes": "Bulk clearance"})
    elif target == 'REORDER':
        tgt = random.randint(1, max(1, reorder))
        if stock > tgt:
            item_txns.append({"Date": fd, "Item ID": iid, "Type": "OUT", "Quantity": stock - tgt,
                              "Reference": ref("OUT"), "Notes": random.choice(NOTES["OUT"])})
        elif stock < tgt:
            item_txns.append({"Date": fd, "Item ID": iid, "Type": "IN", "Quantity": tgt - stock,
                              "Reference": ref("IN"), "Notes": random.choice(NOTES["IN"])})
    elif target == 'OK':
        floor = reorder + 5
        if stock < floor:
            item_txns.append({"Date": fd, "Item ID": iid, "Type": "IN",
                              "Quantity": floor - stock + random.randint(10, 50),
                              "Reference": ref("IN"), "Notes": "Replenishment"})
    all_txns.extend(item_txns)

all_txns.sort(key=lambda t: t['Date'])
for n, t in enumerate(all_txns, start=1):
    t['Transaction ID'] = f"TXN-{n:05d}"

# ─── Write data ─────────────────────────────────────────────────────────────
ws = wb['Item_Master']
item_cols = ["Item ID", "Item Name", "Category", "Brand", "Unit", "Supplier ID",
             "Cost Price", "Selling Price", "Reorder Level", "Opening Stock"]
ws.append(item_cols)
for it in items:
    ws.append([it[c] for c in item_cols])

ws = wb['Suppliers']
sup_cols = ["Supplier ID", "Supplier Name", "City", "State", "Contact Person", "Phone", "Payment Terms"]
ws.append(sup_cols)
for s in suppliers:
    ws.append([s[c] for c in sup_cols])

ws = wb['Stock_Transactions']
tx_cols = ["Transaction ID", "Date", "Item ID", "Type", "Quantity", "Reference", "Notes"]
ws.append(tx_cols)
for t in all_txns:
    ws.append([t[c] for c in tx_cols])

ws = wb['Current_Stock']
cs_cols = ["Item ID", "Item Name", "Category", "Opening Stock",
           "Total Purchased", "Total Sold", "Damaged Qty", "Returned Qty",
           "Current Stock", "Unit Cost", "Stock Value", "Reorder Level", "Stock Status"]
ws.append(cs_cols)
for it in items:
    ws.append([it["Item ID"]])  # Item ID only; rest are header-only

# ─── Styling: lime headers on the four populated sheets ─────────────────────
header_font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill('solid', start_color='5B8C00')
header_align = Alignment(horizontal='center', vertical='center')
thin = Side(border_style='thin', color='BFBFBF')
header_border = Border(left=thin, right=thin, top=thin, bottom=thin)

for name in ['Item_Master', 'Suppliers', 'Stock_Transactions', 'Current_Stock']:
    ws = wb[name]
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'
    for col_cells in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 50)):
        max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 32)

wb.save(OUT_PATH)
print(f"Built: {len(items)} items, {len(suppliers)} suppliers, {len(all_txns)} transactions")
